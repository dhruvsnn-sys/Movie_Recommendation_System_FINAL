"""
Build the final submission documents after the analysis has been run.

Commands:
    python movie_recommender.py
    python build_submission.py

Creates:
    Final_Project_Report.docx
    Final_Project_Presentation.pptx
    Final_Project_Results.md
"""

from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
RESULTS = OUT / "project_results.json"


def load_results():
    if not RESULTS.exists():
        raise SystemExit(
            "Run movie_recommender.py first. It creates outputs/project_results.json."
        )
    return json.loads(RESULTS.read_text(encoding="utf-8"))


def fmt(x, digits=4):
    return f"{x:.{digits}f}" if isinstance(x, (int, float)) else str(x)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_picture_with_caption(doc, filename, caption):
    path = OUT / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.2))
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.italic = True
            run.font.size = Pt(9)


def build_report(r):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MOVIE RECOMMENDATION SYSTEM")
    run.bold = True
    run.font.size = Pt(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("GYEST305 – Introduction to Artificial Intelligence and Data Science\n").bold = True
    p.add_run("III Semester – B.Tech | Department of Engineering / EEE\n")
    p.add_run("Batch 3 – Mini / Micro Project")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "Team: Anush M; Dhruv S Nambiar; Joel C M; Joseph Christy VS; Madhav K Bhattathiri"
    ).italic = True

    add_heading(doc, "1. Introduction", 1)
    doc.add_paragraph(
        "A recommendation system suggests items to users based on information about "
        "their previous interactions or preferences. This project develops a movie "
        "recommendation system using collaborative filtering. The system learns from "
        "MovieLens user ratings, represents them as a user-item matrix, measures "
        "similarity between movies, predicts preferences for unseen movies, and "
        "generates recommendations."
    )

    add_heading(doc, "1.1 Problem Statement", 2)
    doc.add_paragraph(
        "Develop a system that recommends movies to users based on their previous "
        "ratings or similarities in user/item preferences."
    )

    add_heading(doc, "1.2 Objectives", 2)
    add_bullets(doc, [
        "Understand users, movies and ratings.",
        "Inspect and clean the MovieLens ratings data.",
        "Perform exploratory data analysis using meaningful visualizations.",
        "Construct a user-item rating matrix.",
        "Apply item-item collaborative filtering using cosine similarity.",
        "Generate personalized and similar-movie recommendations.",
        "Evaluate rating prediction and top-10 recommendation quality.",
        "Explain sparse data and the cold-start problem.",
    ])

    add_heading(doc, "1.3 Scope", 2)
    doc.add_paragraph(
        "The project is an educational recommender-system implementation using the "
        "MovieLens latest-small development dataset. It demonstrates the complete "
        "data-science workflow and is not intended to be a production recommender."
    )

    add_heading(doc, "2. Dataset", 1)
    d = r["dataset"]
    doc.add_paragraph(
        f"The project uses the MovieLens latest-small dataset from GroupLens. "
        f"The dataset used by this run contains {d['ratings']:,} ratings from "
        f"{d['users']:,} users, with {d['movies_in_catalog']:,} movies in the movie "
        f"catalog and {d['movies_with_ratings']:,} movies represented in the ratings."
    )
    doc.add_paragraph(
        "The ratings file contains userId, movieId, rating and timestamp. "
        "The movie file contains movieId, title and genres. Ratings are on a 0.5–5.0 "
        "five-star scale."
    )

    add_heading(doc, "2.1 Dataset Source", 2)
    doc.add_paragraph(
        "GroupLens Research, MovieLens latest-small dataset: "
        "https://files.grouplens.org/datasets/movielens/ml-latest-small-README.html"
    )
    doc.add_paragraph(
        "Dataset citation: F. Maxwell Harper and Joseph A. Konstan. 2015. "
        "The MovieLens Datasets: History and Context. ACM Transactions on "
        "Interactive Intelligent Systems 5(4), Article 19."
    )

    add_heading(doc, "2.2 Dataset Understanding", 2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Field", "Meaning", "Role"
    for a,b,c in [
        ("userId", "Anonymized user identifier", "User"),
        ("movieId", "Movie identifier", "Item"),
        ("rating", "0.5–5.0 user rating", "Interaction/target"),
        ("timestamp", "Rating time", "Ordering/evaluation"),
        ("title", "Movie title", "Metadata"),
        ("genres", "Pipe-separated genres", "Metadata"),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = a,b,c

    add_heading(doc, "3. Data Preprocessing", 1)
    q = r["quality"]
    doc.add_paragraph(
        "The program explicitly checks missing values, duplicate records, data types, "
        "rating range, and row counts before modelling. Rows with missing required "
        "fields or invalid ratings are removed; duplicate records are removed."
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Quality check"
    table.rows[0].cells[1].text = "Observed result"
    checks = [
        ("Missing cells in ratings", q["ratings_missing_cells"]),
        ("Missing cells in movies", q["movies_missing_cells"]),
        ("Duplicate rating rows", q["ratings_duplicate_rows"]),
        ("Duplicate movie rows", q["movies_duplicate_rows"]),
        ("Invalid rating rows", q["invalid_rating_rows"]),
        ("Rows after cleaning", q["rows_after_cleaning"]),
        ("Rows removed", q["rows_removed_by_cleaning"]),
    ]
    for a,b in checks:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = str(a), str(b)

    add_heading(doc, "3.1 User-Item Matrix", 2)
    doc.add_paragraph(
        "A pivot table is used to construct the user-item matrix. Rows represent users, "
        "columns represent movies, and a cell contains a rating when the user has rated "
        "that movie. Unrated combinations remain missing. This matrix is sparse because "
        "each user rates only a small fraction of the complete movie catalog."
    )

    add_heading(doc, "4. Exploratory Data Analysis", 1)
    doc.add_paragraph(
        "The EDA examines the distribution of ratings, how many ratings users provide, "
        "which movies receive many ratings, and the sparsity pattern of the user-item matrix."
    )

    add_picture_with_caption(
        doc, "01_rating_distribution.png",
        "Figure 1. Rating distribution. This shows how frequently each half-star rating occurs."
    )
    add_picture_with_caption(
        doc, "02_ratings_per_user.png",
        "Figure 2. Number of ratings per user. This shows the variation in user activity."
    )
    add_picture_with_caption(
        doc, "03_ratings_per_movie.png",
        "Figure 3. Top 15 movies by number of ratings. This shows that rating activity is concentrated on a subset of movies."
    )
    add_picture_with_caption(
        doc, "04_user_item_heatmap.png",
        "Figure 4. Sample user-item matrix. Empty cells demonstrate the sparsity of observed ratings."
    )

    add_heading(doc, "5. Methodology", 1)
    doc.add_paragraph(
        "The project uses item-item collaborative filtering. The method does not use "
        "movie plot descriptions to determine similarity. Instead, two movies are considered "
        "similar when their user-rating patterns are similar."
    )

    add_heading(doc, "5.1 Cosine Similarity", 2)
    doc.add_paragraph(
        "For two movie rating vectors A and B, cosine similarity is conceptually "
        "cos(A,B) = (A·B) / (||A|| ||B||). A value closer to 1 indicates that the "
        "vectors point in a similar direction."
    )

    add_heading(doc, "5.2 Rating Prediction", 2)
    doc.add_paragraph(
        "For an unseen movie, the system uses the user's ratings of other movies and "
        "weights those ratings by the similarity between the candidate movie and each "
        "already-rated movie. The weighted average is clipped to the 0.5–5.0 rating range."
    )

    add_heading(doc, "5.3 Recommendation Generation", 2)
    add_bullets(doc, [
        "Select a user and identify movies already rated by that user.",
        "Exclude movies already seen by the user.",
        "Calculate a predicted rating for each candidate movie.",
        "Sort candidates by predicted rating.",
        "Return the top N movies.",
    ])

    add_picture_with_caption(
        doc, "05_similarity_toy_story.png",
        "Figure 5. Example similarity visualization for Toy Story."
    )

    add_heading(doc, "6. Evaluation", 1)
    e = r["evaluation"]
    doc.add_paragraph(
        "To evaluate rating prediction without using the held-out rating to construct "
        "the similarity model, the latest rating of each user is held out when the user "
        "has multiple ratings. Similarity is calculated using the remaining training ratings."
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Result"
    for a,b in [
        ("Training ratings", f"{e['train_ratings']:,}"),
        ("Test ratings", f"{e['test_ratings']:,}"),
        ("MAE", fmt(e["mae"])),
        ("RMSE", fmt(e["rmse"])),
        ("Hit Rate@10", fmt(e["hit_rate_at_10"])),
        ("Precision@10", fmt(e["precision_at_10"])),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = a,b

    doc.add_paragraph(
        "MAE represents the average absolute difference between predicted and actual "
        "ratings. RMSE gives greater weight to larger prediction errors. Hit Rate@10 "
        "records the proportion of relevant held-out user items that appear in the "
        "top-10 list. Precision@10 measures the fraction of the top-10 recommendations "
        "that correspond to the held-out relevant item under this one-item-per-user test."
    )

    add_heading(doc, "7. Results and Interpretation", 1)
    doc.add_paragraph(
        f"The evaluation produced an MAE of {e['mae']:.4f} and an RMSE of {e['rmse']:.4f}. "
        f"The top-10 evaluation produced a Hit Rate@10 of {e['hit_rate_at_10']:.4f} "
        f"and Precision@10 of {e['precision_at_10']:.4f}. These values should be "
        "interpreted as results for this specific dataset, split strategy and similarity "
        "method; they are not a universal benchmark for recommender systems."
    )

    add_heading(doc, "7.1 Example Recommendation Output", 2)
    recs = r["recommendations"]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Movie"
    table.rows[0].cells[1].text = "Predicted rating"
    table.rows[0].cells[2].text = "Genres"
    for item in recs[:10]:
        cells = table.add_row().cells
        cells[0].text = str(item["title"])
        cells[1].text = fmt(float(item["predicted_rating"]), 3)
        cells[2].text = str(item.get("genres", ""))

    add_heading(doc, "8. Discussion", 1)
    doc.add_paragraph(
        "The analysis shows how user-item interactions can be transformed into a "
        "similarity-based recommender. The approach is transparent: recommendations "
        "can be traced to movies that the user has already rated and to the similarity "
        "between those movies and candidate items. The evaluation also demonstrates "
        "that recommendation quality depends on the train/test strategy, data sparsity, "
        "and the similarity method."
    )
    doc.add_paragraph(
        "The method does not use movie content such as plot descriptions. Therefore, "
        "two movies with similar themes may not be considered similar unless their "
        "rating patterns are also similar."
    )

    add_heading(doc, "9. Conclusion", 1)
    doc.add_paragraph(
        "The project implements a complete introductory collaborative-filtering "
        "workflow: dataset inspection, cleaning checks, EDA, user-item matrix creation, "
        "cosine similarity, rating prediction, recommendation generation, evaluation, "
        "interpretation and cold-start handling."
    )

    add_heading(doc, "10. Limitations", 1)
    add_bullets(doc, [
        "The MovieLens latest-small dataset is intended for education and development and can change over time.",
        "The user-item matrix is sparse.",
        "New users have little or no personal history for personalization.",
        "The method uses rating patterns rather than movie semantic/content features.",
        "The evaluation uses one held-out latest rating per eligible user, so it is not the only possible evaluation protocol.",
        "The project is a small educational recommender and does not include production-scale ranking, monitoring or online learning.",
    ])

    add_heading(doc, "11. Future Scope", 1)
    add_bullets(doc, [
        "Use a stable benchmark dataset for reproducible model comparison.",
        "Add matrix factorization such as SVD.",
        "Develop a hybrid recommender using genres/content plus collaborative filtering.",
        "Add more ranking metrics such as Recall@K, NDCG@K and MAP@K.",
        "Add a user interface for collecting new ratings.",
        "Add movie posters and richer metadata.",
        "Deploy the recommender as a web service.",
    ])

    add_heading(doc, "12. References", 1)
    refs = [
        "GroupLens Research. MovieLens datasets. https://grouplens.org/datasets/movielens/",
        "GroupLens Research. MovieLens latest-small README. https://files.grouplens.org/datasets/movielens/ml-latest-small-README.html",
        "Harper, F. M., & Konstan, J. A. (2015). The MovieLens Datasets: History and Context. ACM Transactions on Interactive Intelligent Systems, 5(4), 19:1–19:19.",
        "Pedregosa et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
    ]
    for ref in refs:
        doc.add_paragraph(ref, style="List Number")

    add_heading(doc, "Appendix A — Viva Preparation", 1)
    viva = [
        ("What is collaborative filtering?", "A recommendation approach that uses patterns in user-item interactions to recommend items."),
        ("What is a user-item matrix?", "A matrix whose rows are users, columns are items, and cells contain observed interactions such as ratings."),
        ("How is similarity calculated?", "Here, cosine similarity compares movie rating vectors."),
        ("Content-based vs collaborative filtering?", "Content-based filtering uses item features; collaborative filtering uses interaction patterns."),
        ("What is cold start?", "The difficulty of making personalized recommendations when a new user or item has little interaction history."),
        ("Why is the matrix sparse?", "Most users rate only a small fraction of all movies, leaving many user-movie combinations unobserved."),
        ("Why use a temporal holdout?", "It evaluates on a later interaction while avoiding direct use of that rating in model construction."),
        ("What does MAE mean?", "Mean Absolute Error: the average absolute difference between predicted and actual ratings."),
        ("What does RMSE mean?", "Root Mean Squared Error: the square root of mean squared prediction error, giving more weight to large errors."),
        ("Why can results vary?", "Results depend on the dataset version, train/test split, similarity definition and recommendation settings."),
    ]
    for question, answer in viva:
        p = doc.add_paragraph()
        p.add_run(question + " ").bold = True
        p.add_run(answer)

    path = ROOT / "Final_Project_Report.docx"
    doc.save(path)
    return path


def add_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = PPt(20)
    return slide


def build_ppt(r):
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)

    # 1
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Movie Recommendation System"
    slide.placeholders[1].text = (
        "GYEST305 – Introduction to Artificial Intelligence and Data Science\n"
        "Batch 3 | Collaborative Filtering\n"
        "Anush M • Dhruv S Nambiar • Joel C M • Joseph Christy VS • Madhav K Bhattathiri"
    )

    add_slide(prs, "2. Problem Statement & Objectives", [
        "Problem: recommend movies from previous ratings and user/item preference similarity.",
        "Build a user-item matrix and apply collaborative filtering.",
        "Generate personalized and similar-movie recommendations.",
        "Evaluate predictions and top-10 recommendation quality.",
        "Explain sparse data and the cold-start problem.",
    ])

    add_slide(prs, "3. Dataset", [
        f"MovieLens latest-small development dataset.",
        f"{r['dataset']['ratings']:,} ratings | {r['dataset']['users']:,} users | {r['dataset']['movies_in_catalog']:,} movies.",
        "Ratings use a 0.5–5.0 five-star scale.",
        "Main fields: userId, movieId, rating, timestamp, title, genres.",
        "Source: GroupLens Research.",
    ])

    add_slide(prs, "4. Data Preprocessing & Quality", [
        f"Missing rating cells: {r['quality']['ratings_missing_cells']}.",
        f"Duplicate rating rows: {r['quality']['ratings_duplicate_rows']}.",
        f"Invalid rating rows: {r['quality']['invalid_rating_rows']}.",
        f"Rows after cleaning: {r['quality']['rows_after_cleaning']:,}.",
        "Checks include data types, missing values, duplicates and rating range.",
    ])

    add_slide(prs, "5. Exploratory Data Analysis", [
        "Rating distribution.",
        "Number of ratings per user.",
        "Number of ratings per movie.",
        "Sample user-item matrix heatmap.",
        "Interpretation: ratings are unevenly distributed and the matrix is sparse.",
    ])

    add_slide(prs, "6. Methodology", [
        "Create the user-item rating matrix.",
        "Transpose it to obtain movie rating vectors.",
        "Calculate item-item cosine similarity.",
        "Predict unseen ratings using similarity-weighted ratings.",
        "Rank unseen movies and return the top recommendations.",
    ])

    add_slide(prs, "7. Evaluation Method", [
        "Latest rating of each eligible user is held out as test data.",
        "Similarity is learned only from the remaining training ratings.",
        f"MAE = {r['evaluation']['mae']:.4f}.",
        f"RMSE = {r['evaluation']['rmse']:.4f}.",
        f"Hit Rate@10 = {r['evaluation']['hit_rate_at_10']:.4f}; Precision@10 = {r['evaluation']['precision_at_10']:.4f}.",
    ])

    add_slide(prs, "8. Results & Example Output", [
        f"Example user: {r['demo_user']}.",
        "Top recommendations are generated from similarity-weighted ratings.",
        "Toy Story similarity is visualized to demonstrate item-item relationships.",
        "Results are specific to this dataset, split and method.",
    ])

    add_slide(prs, "9. Limitations & Future Scope", [
        "Sparse user-item data and cold-start users limit personalization.",
        "Method does not use plot/content semantics.",
        "Dataset is intended for education/development.",
        "Future: matrix factorization, hybrid filtering and richer ranking metrics.",
    ])

    add_slide(prs, "10. Conclusion", [
        "The project follows the data-science workflow from data inspection to interpretation.",
        "Collaborative filtering provides transparent, similarity-based recommendations.",
        "Evaluation demonstrates how prediction and ranking quality can be measured.",
        "The system provides a foundation for more advanced recommender systems.",
    ])

    path = ROOT / "Final_Project_Presentation.pptx"
    prs.save(path)
    return path


def main():
    r = load_results()
    report = build_report(r)
    ppt = build_ppt(r)

    results_md = ROOT / "Final_Project_Results.md"
    results_md.write_text(
        "# Final Project Results\n\n"
        f"- Users: {r['dataset']['users']:,}\n"
        f"- Movies in catalog: {r['dataset']['movies_in_catalog']:,}\n"
        f"- Ratings: {r['dataset']['ratings']:,}\n"
        f"- MAE: {r['evaluation']['mae']:.4f}\n"
        f"- RMSE: {r['evaluation']['rmse']:.4f}\n"
        f"- Hit Rate@10: {r['evaluation']['hit_rate_at_10']:.4f}\n"
        f"- Precision@10: {r['evaluation']['precision_at_10']:.4f}\n\n"
        "These values are generated by movie_recommender.py from the downloaded dataset.\n",
        encoding="utf-8"
    )
    print(report)
    print(ppt)


if __name__ == "__main__":
    main()
